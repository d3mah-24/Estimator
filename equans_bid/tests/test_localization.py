"""Localization: formatting, translation completeness, and locale-independent numbers."""

import csv
import io
import json
import re
import unittest
from pathlib import Path
from typing import ClassVar

import frappe
from babel.messages.pofile import read_po
from frappe.tests import IntegrationTestCase

from equans_bid import localization as L
from equans_bid import reasons as R
from equans_bid.api import export_tender, get_tender_summary, process_tender_pdf
from equans_bid.commands.seed_data import run as seed
from equans_bid.export import AMOUNT_COLUMNS, BUILD_UP, MATCHED_BY, TEXT_COLUMNS, XLSX_COLUMNS

APP = Path(__file__).resolve().parent.parent
UI_LOCALES = list(L.UI_LOCALES)
# Raw direct-cost column: after the text columns and the formatted/raw pairs (custom columns follow).
RAW_DIRECT = len(TEXT_COLUMNS) + 2 * len(AMOUNT_COLUMNS) - 1
NBSP, NNBSP, RSQUO = "\u00a0", "\u202f", "\u2019"


def identity(msg, lang=None, context=None):
	return msg


class TestFormatting(unittest.TestCase):
	# Exact babel 2.16 output (verified, not guessed). Note: it differs from the browser's Intl
	# (de-CH U+2019 vs ', fr-CH decimal comma vs dot, en no space after CHF): different CLDR data.
	EXPECTED: ClassVar[dict] = {
		"de-CH": f"CHF{NBSP}414{RSQUO}489.12",
		"fr-CH": f"414{NNBSP}489,12{NBSP}CHF",
		"en": "CHF414,489.12",
	}

	def test_chf_414489_12(self):
		for ui, expected in self.EXPECTED.items():
			with self.subTest(ui=ui):
				self.assertEqual(L.Localizer(ui, identity).chf(414489.12), expected)

	def test_missing_price_is_a_dash_in_every_locale(self):
		for ui in UI_LOCALES:
			self.assertEqual(L.Localizer(ui, identity).chf(None), "—")

	def test_unknown_locale_falls_back_to_de_ch(self):
		self.assertEqual(L.resolve("it-CH"), ("de-CH", "de", "de_CH"))


def _placeholders(text: str) -> set:
	return set(re.findall(r"\{(\w*)\}", text))


class TestTranslationFiles(unittest.TestCase):
	"""Every string the app translates has a German and a French entry with the same placeholders."""

	@classmethod
	def setUpClass(cls):
		cls.catalogs = {}
		for lang in ("de", "fr"):
			with open(APP / "locale" / f"{lang}.po", "rb") as fh:
				cat = read_po(fh)
			cls.catalogs[lang] = {(m.context, m.id): m.string for m in cat if m.id}

	def required(self) -> set:
		need = {(None, s) for s in [*R.REASONS.values(), *R.WARNINGS.values()]}
		need |= {(L.CTX_STATUS, s) for s in L.STATUS_LABELS.values()}
		need |= {(L.CTX_UNIT, s) for s in L.UNIT_LABELS.values()}
		need |= {(L.CTX_SCENARIO, s) for s in L.SCENARIO_LABELS.values()}
		need |= {(L.CTX_PRICE_SOURCE, s) for s in L.PRICE_SOURCE_LABELS.values()}
		export_src = (APP / "export.py").read_text()
		headers = {lbl for _, lbl in TEXT_COLUMNS + AMOUNT_COLUMNS + BUILD_UP}
		headers |= {lbl for _, lbl, _ in XLSX_COLUMNS} | set(MATCHED_BY.values())
		headers |= set(re.findall(r"""\bh\(\s*["']([^"']+)["']\s*\)""", export_src))
		need |= {(L.CTX_EXPORT, s) for s in headers}
		# frappe.throw(_("…")) messages in the API and the DocType controllers
		sources = [APP / "api.py", *(APP / "bid_intelligence" / "doctype").glob("*/*.py")]
		for src in sources:
			need |= {(None, s) for s in re.findall(r'_\("([^"]+)"\)', src.read_text())}
		return need

	def test_every_string_translated(self):
		need = self.required()
		self.assertGreater(len(need), 100)
		for lang, catalog in self.catalogs.items():
			missing = sorted(f"{ctx or ''}|{msg}" for ctx, msg in need if not catalog.get((ctx, msg)))
			self.assertEqual(missing, [], f"missing in {lang}.po")

	def test_placeholders_preserved(self):
		for lang, catalog in self.catalogs.items():
			for (_ctx, msgid), msgstr in catalog.items():
				with self.subTest(lang=lang, msgid=msgid):
					self.assertEqual(_placeholders(msgstr), _placeholders(msgid))

	def test_product_family_labels_are_master_data(self):
		"""Family labels live in data/trades/*/families.json, in every UI locale (not in .po)."""
		from equans_bid.trades import available_trades, load_trade

		for folder in available_trades():
			for family in load_trade(folder).families:
				for ui in UI_LOCALES:
					self.assertTrue(family.labels.get(ui), (folder, family.code, ui))


def _rows(content: str) -> list[list[str]]:
	return list(csv.reader(io.StringIO(content.lstrip("﻿")), delimiter=";"))


class TestLocaleIndependence(IntegrationTestCase):
	"""Switching the language changes text and formatting, never a calculated value."""

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		seed()
		cls.tender = frappe.get_doc({"doctype": "Tender", "project_name": "L10N test tender"}).insert()
		process_tender_pdf(cls.tender.name, pages=[78, 79, 80])

	def tearDown(self):
		frappe.local.lang = "en"

	def summary_in(self, lang):
		frappe.local.lang = lang
		frappe.form_dict._lang = lang
		try:
			return get_tender_summary(self.tender.name)
		finally:
			frappe.form_dict.pop("_lang", None)

	def test_totals_identical_in_every_language(self):
		summaries = {lang: self.summary_in(lang) for lang in ("de", "fr", "en")}
		for key in ("totals", "build_up", "bid_range", "indicative_price", "by_status"):
			self.assertEqual(summaries["de"][key], summaries["fr"][key], key)
			self.assertEqual(summaries["de"][key], summaries["en"][key], key)

	def export_rows(self, ui):
		export_tender(self.tender.name, lang=ui, format="csv")
		return _rows(frappe.local.response.filecontent)

	def workbook(self, ui):
		from openpyxl import load_workbook

		export_tender(self.tender.name, lang=ui)
		return load_workbook(io.BytesIO(frappe.local.response.filecontent))

	def test_excel_amounts_are_numbers_and_identical_in_every_locale(self):
		values = {}
		for ui in UI_LOCALES:
			ws = self.workbook(ui).worksheets[0]
			header = [c.value for c in ws[4]]
			ref, direct = header.index(L.Localizer(ui).t("Reference LV", L.CTX_EXPORT)), len(XLSX_COLUMNS) - 1
			rows = list(ws.iter_rows(min_row=5, values_only=True))
			lv1 = next(r for r in rows if r[ref] == "LV0001")
			self.assertEqual(lv1[direct], 806.26)  # a number, not text
			values[ui] = [r[direct] for r in rows]
		self.assertEqual(values["de-CH"], values["fr-CH"])
		self.assertEqual(values["de-CH"], values["en"])

	def test_excel_missing_price_is_dash_and_sheets_translated(self):
		wb = self.workbook("de-CH")
		self.assertEqual(
			wb.sheetnames, ["Positionen", "Zusammenfassung", "Angebotsbereich", "Offene Punkte", "Annahmen"]
		)
		self.assertEqual(self.workbook("fr-CH").sheetnames[0], "Positions")
		ws = wb.worksheets[0]
		doc = frappe.get_doc("Tender", self.tender.name)
		unpriced = {i.position_id for i in doc.items if i.status == "review_required"}
		self.assertTrue(unpriced)
		for r in ws.iter_rows(min_row=5, values_only=True):
			if r[0] in unpriced:
				self.assertEqual(r[len(XLSX_COLUMNS) - 1], "—", r[0])
		open_ids = {r[0] for r in wb["Offene Punkte"].iter_rows(min_row=2, values_only=True)}
		self.assertLessEqual(unpriced, open_ids)

	def test_lv0001_golden_806_26_in_every_locale(self):
		for ui in UI_LOCALES:
			with self.subTest(ui=ui):
				frappe.local.lang = L.resolve(ui)[1]
				doc = frappe.get_doc("Tender", self.tender.name)
				lv1 = next(i for i in doc.items if i.golden_lv_id == "LV0001")
				self.assertAlmostEqual(lv1.direct_cost, 806.26, places=2)
				rows = self.export_rows(ui)
				row = next(r for r in rows if len(r) > 1 and r[1] == "LV0001")
				self.assertEqual(row[RAW_DIRECT], "806.26")  # raw direct-cost column, dot decimal
				self.assertIn("806", row[RAW_DIRECT - 1])  # formatted column, locale separators

	def test_raw_amounts_identical_and_text_translated(self):
		by_locale = {ui: self.export_rows(ui) for ui in UI_LOCALES}
		n_text = len(TEXT_COLUMNS)
		raw_cols = [n_text + 2 * i + 1 for i in range(len(AMOUNT_COLUMNS))]

		def raw_part(rows):
			return [[r[c] for c in raw_cols] for r in rows[3:] if len(r) > max(raw_cols)]

		self.assertEqual(raw_part(by_locale["de-CH"]), raw_part(by_locale["fr-CH"]))
		self.assertEqual(raw_part(by_locale["de-CH"]), raw_part(by_locale["en"]))
		self.assertEqual(by_locale["de-CH"][2][0], "LV-ID")
		self.assertEqual(by_locale["fr-CH"][2][0], "N° de position")
		self.assertEqual(by_locale["en"][2][0], "LV ID")

	def test_missing_price_exports_as_dash_not_zero(self):
		for ui in UI_LOCALES:
			rows = self.export_rows(ui)
			doc = frappe.get_doc("Tender", self.tender.name)
			unpriced = {i.position_id for i in doc.items if i.status == "review_required"}
			self.assertTrue(unpriced)
			for r in rows:
				if r and r[0] in unpriced:
					self.assertEqual(r[RAW_DIRECT - 1], "—", (ui, r[0]))  # formatted direct cost
					self.assertEqual(r[RAW_DIRECT], "", (ui, r[0]))  # raw direct cost

	def test_database_keeps_codes(self):
		for ui in UI_LOCALES:
			self.export_rows(ui)
		statuses = set(frappe.get_all("Tender Item", {"parent": self.tender.name}, pluck="status"))
		self.assertLessEqual(statuses, set(L.STATUS_LABELS))
		units = set(frappe.get_all("Tender Item", {"parent": self.tender.name}, pluck="unit"))
		self.assertLessEqual(units, set(L.UNIT_LABELS))
		for row in frappe.get_all(
			"Tender Item", {"parent": self.tender.name}, ["reason_code", "reason_args"]
		):
			self.assertIn(row.reason_code, R.REASONS)
			json.loads(row.reason_args or "{}")
