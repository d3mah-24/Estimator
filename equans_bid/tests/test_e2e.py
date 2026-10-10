"""End to end in Frappe: PDF -> Tender Items -> price -> estimator edit -> reprice."""

import io

import frappe
from frappe.tests import IntegrationTestCase

from equans_bid.api import (
	export_tender,
	get_tender_summary,
	process_tender_pdf,
	set_bid_assumptions,
	set_bid_selection,
	update_item,
)
from equans_bid.commands.seed_data import run as seed


class TestTenderEndToEnd(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		seed()
		cls.tender = frappe.get_doc({"doctype": "Tender", "project_name": "E2E test tender"}).insert()
		process_tender_pdf(cls.tender.name, pages=[78, 79, 80])

	def item(self, lv_id):
		doc = frappe.get_doc("Tender", self.tender.name)
		return doc, next(i for i in doc.items if i.golden_lv_id == lv_id)

	def test_lv0001_extracted_and_priced(self):
		_, lv1 = self.item("LV0001")
		self.assertEqual(lv1.source_page, 78)
		self.assertEqual(lv1.quantity, 5)
		self.assertEqual(lv1.unit, "pc")
		self.assertEqual(lv1.matched_product, "P-AIR-001")
		self.assertEqual(lv1.status, "calculated")
		self.assertAlmostEqual(lv1.direct_cost, 806.26, places=2)

	def test_supplier_quote_override_reprices(self):
		doc, lv6 = self.item("LV0006")
		self.assertEqual(lv6.status, "supplier_quote_required")
		labour_before = lv6.labour_cost
		update_item(doc.name, lv6.name, quote_override=9000)
		_, lv6 = self.item("LV0006")
		self.assertEqual(lv6.material_cost, 9000)
		self.assertEqual(lv6.labour_cost, labour_before)
		self.assertEqual(lv6.price_source, "supplier_quote")

	def test_reprocessing_keeps_estimator_edits(self):
		doc, lv6 = self.item("LV0006")
		update_item(doc.name, lv6.name, quote_override=7777)
		other = next(i for i in doc.items if i.status == "review_required")
		update_item(doc.name, other.name, status="out_of_scope")
		process_tender_pdf(self.tender.name, pages=[78, 79, 80])
		_, lv6 = self.item("LV0006")
		self.assertEqual((lv6.quote_override, lv6.material_cost), (7777, 7777))
		again = next(
			i for i in frappe.get_doc("Tender", doc.name).items if i.position_id == other.position_id
		)
		self.assertEqual((again.status, again.reason_code), ("out_of_scope", "status_set_by_estimator"))
		update_item(doc.name, lv6.name, quote_override=0)
		update_item(doc.name, again.name, status="review_required")

	def test_bid_selection_saved_and_exported(self):
		s = set_bid_selection(self.tender.name, "Conservative")
		self.assertEqual(s["selected_bid"]["scenario"], "Conservative")
		self.assertTrue(next(x for x in s["bid_range"] if x["scenario"] == "Conservative")["selected"])
		s = set_bid_selection(self.tender.name, "custom", price_chf="500000", margin_pct="")
		self.assertEqual(s["selected_bid"]["bid_price_chf"], 500000)
		self.assertFalse(any(x["selected"] for x in s["bid_range"]))

		from openpyxl import load_workbook

		export_tender(self.tender.name)
		wb = load_workbook(io.BytesIO(frappe.local.response.filecontent))
		labels = {r[0]: r[1:] for r in wb["Zusammenfassung"].iter_rows(values_only=True) if r and r[0]}
		self.assertEqual(labels["Gewähltes Angebot"][0], 500000)
		bids = [r for r in wb["Angebotsbereich"].iter_rows(min_row=2, values_only=True) if r[0]]
		self.assertIn(("Eigenes Angebot", 500000), [(r[0], r[1]) for r in bids])

		with self.assertRaises(frappe.ValidationError):
			set_bid_selection(self.tender.name, "custom", margin_pct=50)
		self.assertIsNone(set_bid_selection(self.tender.name, None)["selected_bid"])

	def test_bid_assumptions_are_saved_and_reset(self):
		before = get_tender_summary(self.tender.name)
		after = set_bid_assumptions(self.tender.name, {"scenario_margins_pct": {"Target": 8}})
		self.assertEqual(after["assumptions"]["scenario_margins_pct"]["Target"], 8)
		target = next(s for s in after["bid_range"] if s["scenario"] == "Target")
		old = next(s for s in before["bid_range"] if s["scenario"] == "Target")
		self.assertGreater(target["bid_price_chf"], old["bid_price_chf"])
		reset = set_bid_assumptions(self.tender.name, None)
		self.assertEqual(reset["assumptions"], reset["default_assumptions"])

	def test_export_excel(self):
		from openpyxl import load_workbook

		export_tender(self.tender.name)
		self.assertTrue(frappe.local.response.filename.endswith(".xlsx"))
		wb = load_workbook(io.BytesIO(frappe.local.response.filecontent))
		summary = wb["Zusammenfassung"]
		labels = {r[0]: r[1] for r in summary.iter_rows(values_only=True) if r and r[0]}
		baseline = frappe.get_doc("Tender", self.tender.name).cost_baseline
		self.assertAlmostEqual(labels["Kostenbasis"], baseline, places=2)

	def test_export_csv(self):
		export_tender(self.tender.name, format="csv")
		content = frappe.local.response.filecontent
		self.assertEqual(frappe.local.response.type, "download")
		self.assertTrue(frappe.local.response.filename.endswith(".csv"))
		self.assertIn("LV-ID;Referenz-LV;Anlage;Seite", content)
		self.assertIn("Kostenbasis", content)
		self.assertIn("806.26", content)

	def test_unrecognised_positions_stay_unpriced(self):
		doc = frappe.get_doc("Tender", self.tender.name)
		for i in doc.items:
			if i.status in ("review_required", "insufficient_information"):
				self.assertEqual(i.direct_cost, 0, i.position_id)
		self.assertGreater(doc.unresolved_count, 0)
